import os
import logging
import json
import re
from typing import Dict, List, Tuple, Optional
from src.text_processing import clean_text, detect_language
import asyncio
from src.user_management import *
from src.prompts_summarize import VALIDATION_ERRORS
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.constants import ParseMode
from telegram.ext import ContextTypes

# Logging setup
logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO
)
logger = logging.getLogger(__name__)

QUIZ_TYPES = {
    "mc": "Multiple Choice",
    "oe": "Open Ended"
}


BUTTON_TRANSLATIONS = {
    "en": {
        "start_quiz": "🎯 Start Quiz",
        "share_quiz": "📤 Share Quiz",
        "share_to_group": "👥 Share to Group",
        "download_word": "📥 Download as Word",
        "multiplayer_battle": "🎮 Multiplayer Battle"
    },
    "uz": {
        "start_quiz": "🎯 Quizni boshlash",
        "share_quiz": "📤 Quizni ulashish",
        "share_to_group": "👥 Guruhga ulashish",
        "download_word": "📥 Word formatda yuklash",
        "multiplayer_battle": "🎮 Ko'p o'yinchi jang"
    },
    "ru": {
        "start_quiz": "🎯 Начать квиз",
        "share_quiz": "📤 Поделиться квизом",
        "share_to_group": "👥 Поделиться в группе",
        "download_word": "📥 Скачать как Word",
        "multiplayer_battle": "🎮 Мультиплеер битва"
    }
}

def clean_json_string(json_str: str) -> str:
    import re
    json_str = re.sub(r'```json\s*|\s*```', '', json_str)
    json_str = json_str.strip()
    json_str = re.sub(r',(\s*[}\]])', r'\1', json_str)
    if not json_str.startswith('{'):
        json_str = '{' + json_str
    if not json_str.endswith('}'):
        json_str = json_str + '}'
    return json_str


from html import escape as _html_escape

TG_LIMIT = 4000   # Telegram hard limit is 4096 characters per message

def format_question_message(number, total, question, options=None, max_option_len=100):
    head = f"<b>{number}/{total}</b>\n\n"
    q = question or ""
    opt_lines, shown = "", []
    if options and any(len(o) > max_option_len for o in options):
        shown = [o if len(o) <= 400 else o[:397] + "..." for o in options]
        opt_lines = "\n\n" + "\n".join(
            f"<b>{chr(65 + i)})</b> {_html_escape(o)}" for i, o in enumerate(shown))
    visible_extra = len(f"{number}/{total}\n\n") + (
        len("\n\n" + "\n".join(f"{chr(65 + i)}) {o}" for i, o in enumerate(shown))) if opt_lines else 0)
    room = max(TG_LIMIT - visible_extra, 500)
    if len(q) > room:
        q = q[:room - 3] + "..."
    return f"{head}{_html_escape(q)}{opt_lines}"


async def get_validation_error_message(user_id: str, validation_reason: str) -> str:
    from src.user_management import get_user_language

    language = await get_user_language(user_id)
    lang = language if language in VALIDATION_ERRORS else "en"
    msgs = VALIDATION_ERRORS[lang]
    
    return (
        f"*{msgs['title']}*\n\n"
        f"*{msgs['reason']}*: {validation_reason}\n\n"
        f"{msgs['instruction']}\n\n"
        f"_{msgs['restart']}_"
    )

def truncate_content_for_processing(text: str, max_chars: int = 5000) -> str:
    if len(text) <= max_chars:
        return text
    
    # Try to break at sentence boundary
    truncated = text[:max_chars]
    
    # Look for last sentence-ending punctuation
    last_period = max(
        truncated.rfind('.'),
        truncated.rfind('!'),
        truncated.rfind('?'),
        truncated.rfind('\n')
    )
    
    if last_period > max_chars * 0.7:  # If we found a good break point (at least 70% through)
        return truncated[:last_period + 1].strip()
    
    # Otherwise just truncate at word boundary
    last_space = truncated.rfind(' ')
    if last_space > 0:
        return truncated[:last_space].strip() + "..."
    
    return truncated.strip() + "..."

async def validate_text_with_llm(text: str, user_id: str) -> Tuple[bool, str, Optional[str]]:
    from src.imports import async_generate_content_with_retry
    
    # Truncate before cleaning to avoid processing huge texts
    MAX_VALIDATION_LENGTH = 3000
    if len(text) > MAX_VALIDATION_LENGTH:
        logger.warning(f"[VALIDATION] Text length {len(text)} exceeds max, truncating to {MAX_VALIDATION_LENGTH}")
        text = text[:MAX_VALIDATION_LENGTH] + "..."
    
    cleaned_text = clean_text(text)
    logger.info(f"[VALIDATION] User ID: {user_id}, Cleaned text length: {len(cleaned_text)} characters")

    if not user_id or (isinstance(user_id, str) and not user_id.strip()):
        logger.warning(f"[VALIDATION] Invalid user_id: {user_id}")
        user_id = "0"

    if len(cleaned_text.strip()) < 2:
        logger.warning(f"[VALIDATION] Text too short: {len(cleaned_text.strip())} characters")
        return False, await get_validation_error_message(user_id, "Text is too short to analyze."), None

    # Truncate validation text snippet for the prompt
    validation_text = cleaned_text[:2000] if len(cleaned_text) > 2000 else cleaned_text
    
    validation_prompt = f"""
You are an expert educational content validator. Your task is to determine if the given text contains meaningful educational content that can be used to create quiz questions.

Analyze the following text: "{validation_text}"

Consider these guidelines:
1. EDUCATIONAL TERMS/NAMES are VALID even if short:
   - Historical events (e.g., "fermopil janggi", "Battle of Thermopylae", "World War II")
   - Scientific concepts (e.g., "photosynthesis", "DNA", "gravity")
   - Famous people, places, or concepts
   - Subject-specific terminology

2. NON-EDUCATIONAL content is INVALID:
   - Generic filenames (document.pdf, file.docx)
   - Random characters or numbers
   - Greetings only (hello, hi, salom)
   - Meaningless text or gibberish
   - Pure navigation commands (/start, /help)

3. LANGUAGE CONSIDERATIONS:
   - Educational content in any language is valid
   - Transliterations and local spellings are valid
   - Historical or cultural terms in native languages are valid

4. CONTEXT AWARENESS:
   - Even single educational terms can generate multiple quiz questions
   - Famous battles, concepts, people can have rich quiz content
   - Don't reject based on brevity if educationally meaningful

IMPORTANT: "fermopil janggi" refers to the Battle of Thermopylae, a famous historical battle - this IS educational content and should be VALID.

Respond with ONLY a JSON object:
{{
    "is_valid": true/false,
    "reason": "Brief explanation",
    "language": "detected language code (e.g., 'en', 'uz', 'ru')"
}}
"""
    
    try:
        # Convert user_id to int for proper key assignment
        user_id_int = int(user_id) if isinstance(user_id, str) and user_id.isdigit() else None
        
        # Use the new retry system
        response = await async_generate_content_with_retry(
            validation_prompt, 
            user_id=user_id_int,
            max_attempts=10
        )
        
        if not response:
            logger.error("[VALIDATION] No response from LLM")
            return False, await get_validation_error_message(user_id, "Validation failed - no response from AI"), None

        cleaned_response = clean_json_string(response)
        validation_result = json.loads(cleaned_response)

        is_valid = validation_result.get("is_valid", False)
        reason = validation_result.get("reason", "Unable to validate text")
        detected_language = validation_result.get("language", "en")

        if not is_valid:
            # Fallback regex check
            educational_patterns = [
                r'\b(jang|battle|war|urush)\b',
                r'\b(acid|base|molecule|atom|cell|DNA|RNA)\b',
                r'\b(equation|theorem|formula|algebra|geometry)\b',
                r'\b(tarix|history|science|matematik|biology|chemistry|physics)\b'
            ]
            if any(re.search(p, cleaned_text.lower()) for p in educational_patterns):
                return True, "Educational content detected by pattern matching", detected_language
            return False, await get_validation_error_message(user_id, reason), None

        return True, reason, detected_language

    except json.JSONDecodeError as e:
        logger.error(f"[VALIDATION] JSON decode error: {e}")
        return False, await get_validation_error_message(user_id, "Invalid response format from AI"), None

    except Exception as e:
        logger.error(f"[VALIDATION] Error: {e}")
        return False, await get_validation_error_message(user_id, "Validation failed due to an error"), None


async def generate_quiz(
    text: str, quiz_type: str, detected_language: str, user_id: str, num_questions: int = 5
) -> Tuple[bool, str, Optional[Dict]]:

    from src.user_management import get_user_language
    from src.imports import async_generate_content_with_retry
    from src.prompts_summarize import MCH_QUIZ_PROMPT, OE_QUIZ_PROMPT

    # Truncate before processing
    MAX_QUIZ_GEN_LENGTH = 10000
    if len(text) > MAX_QUIZ_GEN_LENGTH:
        logger.warning(f"[QUIZ_GEN] Text length {len(text)} exceeds max, truncating to {MAX_QUIZ_GEN_LENGTH}")
        text = text[:MAX_QUIZ_GEN_LENGTH] + "..."

    user_language = await get_user_language(user_id)
    cleaned_text = clean_text(text)
    doc_language = detected_language

    # FORMAT THE PROMPTS WITH num_questions AND doc_language
    if quiz_type == "mc":
        system_prompt = MCH_QUIZ_PROMPT.format(
            num_questions=num_questions,
            doc_language=doc_language
        )
    else:
        system_prompt = OE_QUIZ_PROMPT.format(
            num_questions=num_questions,
            doc_language=doc_language
        )

    user_prompt = f"""Text for quiz generation:\n\n{cleaned_text}"""

    full_prompt = f"{system_prompt}\n\n{user_prompt}"

    # JSON parsing retry configuration
    MAX_JSON_RETRIES = 5
    json_retry_count = 0
    last_error = None
    
    try:
        # Convert user_id to int for proper key assignment
        user_id_int = int(user_id) if isinstance(user_id, str) and user_id.isdigit() else None
        
        # Try to generate and parse quiz with retries
        while json_retry_count < MAX_JSON_RETRIES:
            try:
                # Use the same API key for retries
                response = await async_generate_content_with_retry(
                    full_prompt,
                    user_id=user_id_int,
                    max_attempts=10
                )

                if not response:
                    logger.error(f"[QUIZ_GEN] No response from LLM for user {user_id}")
                    json_retry_count += 1
                    last_error = "No response from AI"
                    continue

                logger.info(f"[QUIZ_GEN] LLM response received, length: {len(response)}")
                logger.debug(f"[QUIZ_GEN] Raw response: {response[:500]}")

                # Try to parse JSON
                cleaned_response = clean_json_string(response)
                quiz_data = json.loads(cleaned_response)

                # Enhanced validation
                if not isinstance(quiz_data, dict):
                    raise ValueError("Invalid quiz data structure")
                
                if quiz_type == "mc":
                    if "questions" not in quiz_data or not quiz_data["questions"]:
                        raise ValueError("No MC questions generated")
                    
                    # Validate question quality
                    for q in quiz_data["questions"]:
                        if not q.get("question") or not q.get("options") or not q.get("correct_answer"):
                            raise ValueError("Incomplete question structure")
                        if len(q["options"]) != 4:
                            raise ValueError("Each question must have exactly 4 options")
                        if q["correct_answer"] not in q["options"]:
                            raise ValueError("Correct answer must be one of the options")
                else:
                    if "questions" not in quiz_data or not quiz_data["questions"]:
                        raise ValueError("No OE questions generated")
                    
                    for q in quiz_data["questions"]:
                        if not q.get("question") or not q.get("answer"):
                            raise ValueError("Incomplete question structure")

                logger.info(f"[QUIZ_GEN] Quiz validation passed, returning data")
                return True, "✅ Quiz generated successfully", quiz_data

            except json.JSONDecodeError as e:
                json_retry_count += 1
                last_error = f"JSON parsing error: {str(e)}"
                logger.error(f"[QUIZ_GEN] JSON decode error (attempt {json_retry_count}/{MAX_JSON_RETRIES}): {e}")
                
                if json_retry_count < MAX_JSON_RETRIES:
                    logger.info(f"[QUIZ_GEN] Retrying JSON parse... ({json_retry_count}/{MAX_JSON_RETRIES})")
                    await asyncio.sleep(1)  # Brief delay before retry
                    continue
                else:
                    logger.error(f"[QUIZ_GEN] Max JSON retries reached for user {user_id}")
                    break

            except ValueError as e:
                json_retry_count += 1
                last_error = f"Validation error: {str(e)}"
                logger.error(f"[QUIZ_GEN] Validation error (attempt {json_retry_count}/{MAX_JSON_RETRIES}): {e}")
                
                if json_retry_count < MAX_JSON_RETRIES:
                    logger.info(f"[QUIZ_GEN] Retrying due to validation error... ({json_retry_count}/{MAX_JSON_RETRIES})")
                    await asyncio.sleep(1)
                    continue
                else:
                    logger.error(f"[QUIZ_GEN] Max validation retries reached for user {user_id}")
                    break

        # All retries failed. The caller refunds (it knows what was actually charged).
        logger.error(f"[QUIZ_GEN] All retries failed for user {user_id}: {last_error}")
        return False, (
            "❌ *Quiz Generation Failed*\n\n"
            f"We couldn't generate your quiz after {MAX_JSON_RETRIES} attempts.\n\n"
            f"Please try again with different content.\n\n_Error: {last_error}_"
        ), None

    except Exception as e:
        logger.error(f"[QUIZ_GEN] Unexpected error in quiz generation: {e}")
        return False, f"❌ Error generating quiz: {str(e)}", None


async def generate_quiz_docx(quiz_data: Dict, quiz_type: str) -> io.BytesIO:
    """Word document for teachers: questions first, answer key at the very end."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    questions = quiz_data.get("questions", [])

    if quiz_type == "mc":
        title = doc.add_heading(quiz_data.get("title", "Multiple Choice Quiz"), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        summary = quiz_data.get("summary", "")
        if summary:
            p = doc.add_paragraph(summary)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()

        answer_lines = []
        for i, q in enumerate(questions, 1):
            para = doc.add_paragraph()
            run = para.add_run(f"Question {i}: {q['question']}")
            run.bold = True
            run.font.size = Pt(12)

            options = q.get('options', [])
            for j, opt in enumerate(options):
                op = doc.add_paragraph(f"{chr(65 + j)}. {opt}")
                op.paragraph_format.left_indent = Inches(0.5)

            # collect the answer(s) for the key at the end
            if q.get('allows_multiple_answers') and q.get('correct_answers') is not None:
                idxs = []
                for a in q['correct_answers']:
                    if isinstance(a, int):
                        idxs.append(a)
                    elif a in options:
                        idxs.append(options.index(a))
            else:
                a = q.get('correct_answer')
                if isinstance(a, int):
                    idxs = [a]
                elif a in options:
                    idxs = [options.index(a)]
                else:
                    idxs = []

            idxs = [k for k in idxs if 0 <= k < len(options)]
            if idxs:
                letters = ", ".join(chr(65 + k) for k in idxs)
                texts = "; ".join(options[k] for k in idxs)
                answer_lines.append(f"{i}. {letters} - {texts}")
            else:
                answer_lines.append(f"{i}. N/A")

            doc.add_paragraph()

        # answer key on its own page, at the end
        doc.add_page_break()
        head = doc.add_paragraph()
        hr = head.add_run("Answers:")
        hr.bold = True
        hr.font.size = Pt(14)
        for line in answer_lines:
            ap = doc.add_paragraph()
            num, rest = line.split(". ", 1)
            ap.add_run(f"{num}. ").bold = True
            r = ap.add_run(rest)
            r.font.color.rgb = RGBColor(0, 128, 0)

    else:  # open-ended
        title = doc.add_heading(quiz_data.get("title", "Open-Ended Quiz"), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

        for i, q in enumerate(questions, 1):
            para = doc.add_paragraph()
            run = para.add_run(f"{i}. {q['question']}")
            run.bold = True
            run.font.size = Pt(12)
            doc.add_paragraph()

        doc.add_page_break()
        head = doc.add_paragraph()
        hr = head.add_run("Answers:")
        hr.bold = True
        hr.font.size = Pt(14)
        for i, q in enumerate(questions, 1):
            ap = doc.add_paragraph()
            ap.add_run(f"{i}. ").bold = True
            ap.add_run(str(q.get('answer', 'N/A')))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf



async def handle_download_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE):
    from src.user_management import get_user_language, get_shared_quiz, get_user_quiz_type 
    query = update.callback_query
    await query.answer()
    
    user_id = query.from_user.id
    language = await get_user_language(user_id)
    
    try:
        # Extract quiz_id from callback data
        quiz_id = query.data.split("_", 2)[2]
        
        # Get quiz data
        quiz_data = context.user_data.get('current_quiz')
        if not quiz_data:
            # Try to get from database if not in context
            quiz_data = await get_shared_quiz(quiz_id)
        
        if not quiz_data:
            await query.message.reply_text(
                "❌ Quiz data not found. Please generate a new quiz.",
                parse_mode=ParseMode.MARKDOWN
            )
            return
        
        # Show processing message
        status_msg = await query.message.reply_text(
            "📄 Generating Word document...",
            parse_mode=ParseMode.MARKDOWN
        )
        
        # Generate DOCX
        quiz_type = await get_user_quiz_type(user_id)
        docx_buffer = await generate_quiz_docx(quiz_data, quiz_type)
        
        # Delete status message
        await status_msg.delete()
        
        # Send document
        title = quiz_data.get("title", "quiz").replace("/", "-")
        filename = f"{title[:50]}.docx"
        
        await query.message.reply_document(
            document=docx_buffer,
            filename=filename,
            caption="📥 Your quiz in Word format"
        )
        
    except Exception as e:
        logger.error(f"Error generating Word document: {e}")
        await query.message.reply_text(
            "❌ Error generating document. Please try again.",
            parse_mode=ParseMode.MARKDOWN
        )

def format_quiz_for_telegram(quiz_data: Dict, quiz_type: str) -> List[str]:
    if quiz_type == "mc":
        title = quiz_data.get("title", "📝 *Multiple Choice Quiz*")
        summary = quiz_data.get("summary", "")
        
        formatted_quiz = f"{title}\n\n"
        if summary:
            formatted_quiz += f"*Summary:* {summary}\n\n"
            
        for i, q in enumerate(quiz_data["questions"], 1):
            formatted_quiz += f"<b>Question {i}:</b> {q['question']}\n\n"
            for j, opt in enumerate(q['options'], 1):
                formatted_quiz += f"{j}. {opt}\n"
            formatted_quiz += f"\n<b>Correct Answer:</b> <span class=\"tg-spoiler\">{q['correct_answer']}</span>\n\n"
        return [formatted_quiz]
    else:  # oe
        messages = []
        current_message = "📝 <b>Open-Ended Quiz</b>\n\n"
        for i, q in enumerate(quiz_data["questions"], 1):
            question_text = (
                f"<b>{i}. {q['question']}</b>\n\n"
                f"<b>Answer:</b> <span class=\"tg-spoiler\">{q['answer']}</span>\n\n"
            )
            if len(current_message) + len(question_text) > 4000:
                messages.append(current_message)
                current_message = question_text
            else:
                current_message += question_text
        if current_message:
            messages.append(current_message)
        return messages

def truncate_text(text: str, max_length: int = 99) -> str:
    if len(text) <= max_length:
        return text
    return text[:max_length-3].strip() + "..."

def create_telegram_poll(question: Dict) -> Dict:
    from src.poll_utils import poll_options
    options = question["options"]
    ca = question["correct_answer"]
    correct_option_id = ca if isinstance(ca, int) else options.index(ca)
    return {
        "question": truncate_text(question["question"], 300),
        "options": poll_options(options),
        "is_anonymous": False,
        "type": "quiz",
        "correct_option_id": correct_option_id,
        "explanation": f"Correct answer: {options[correct_option_id]}"[:200],
    }