import logging
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, ReplyKeyboardMarkup, KeyboardButton
from telegram.ext import ContextTypes, ConversationHandler, CommandHandler, MessageHandler, CallbackQueryHandler, filters
from telegram.constants import ParseMode
from handlers.quiz_splitter import QuizSplitter, validate_quiz_structure
import os
from docx import Document
import PyPDF2
import olefile
import asyncio
from datetime import datetime, timezone
import random

logger = logging.getLogger(__name__)


def is_free_user(user_id) -> bool:
    from src.imports import is_free_user as _is_free_user   # lazy: avoids a circular import
    return _is_free_user(user_id)

# Conversation states
SPLIT_WAITING_FILE = 0
SPLIT_CHOOSE_SIZE = 1
SPLIT_CHOOSE_TIME = 2
SPLIT_ASK_QUIZ_INFO = 3 
SPLIT_ASK_RANDOMIZE = 4  
SPLIT_CONFIRM = 5

# Centralized translation dictionary
SPLIT_TRANSLATIONS = {
    "en": {
        "split_quiz_start": "📂 *Split Quiz File*\n\nPlease upload your quiz file (.txt, .docx, .doc, or .pdf format).\n\n*Supported formats:*\n• Questions separated by ==== and ++++\n• Questions with +/- markers\n• Numbered questions (1., 2., 3.)\n• Standard format with #, ✅, or + for correct answers",
        "split_file_received": "✅ *File received!*\n\n📊 Total quizzes found: *{total}*\n\nNow choose how many quizzes per batch:",
        "split_invalid_file": "❌ *Invalid file format* or no valid quizzes found.\n\nPlease make sure:\n• File is .txt, .docx, .doc, or .pdf\n• Questions have at least 2 options\n• Correct answers are marked properly",
        "split_choose_time": "⏱ *Time per Quiz*\n\nChoose time limit for each quiz (in seconds):",
        "split_confirm": "✅ *Confirm Split*\n\n📊 *Summary:*\n• Total quizzes: *{total}*\n• Batch size: *{batch_size}*\n• Number of batches: *{num_batches}*\n• Last batch: *{last_batch}* quizzes\n• Time per quiz: *{time}* seconds",
        "split_success": "✅ *Quiz Split Complete!*\n\n{num_batches} batch(es) created successfully.",
        "split_cancelled": "❌ Quiz split cancelled. Use /start to begin again.",
        "back_to_menu": "↩️ Back to Menu",
        "extracting_text": "📝 Extracting text from {file_type}...",
        "split_ask_quiz_info": "📝 *Quiz Information*\n\nEnter any information about this quiz batch (name, author, channel, etc.).\n\n*Examples:*\n• Sociology Quizzes by Bahtiyor Komilov\n• Philosophy Tests - @SimpleLearnuz\n• History Quiz Collection",
        "split_quiz_info_received": "✅ Quiz info set: *{info}*",
        "quiz_info_too_short": "❌ Quiz info must be at least 3 characters. Please try again:",
        "split_info_invalid": "❌ Please provide both name and author on separate lines.\n\n*Example:*\n```\nQuiz Name\nAuthor Name\n```",

        "batch_name_too_short": "❌ Batch name must be at least 3 characters. Please try again:",
        "author_name_too_short": "❌ Author name must be at least 2 characters. Please try again:",
        "could_not_extract": "❌ Could not extract text from the file.\n\n📋 *Supported Formats:*\n\n*1. Plus/Minus Format:*\n```\n#Question text\n+Correct option\n-Wrong option\n-Wrong option\n```\n\n*2. Numbered Format:*\n```\n1. Question text\nA) #Correct option\nB) Wrong option\nC) Wrong option\n```\n\nPlease ensure your file follows one of these formats.",
        "price_label": "💰 *Price: {price:,} UZS*",
        "doc_error": "❌ Error reading .doc file. Please save as .docx or .txt.\n\nError: {error}",
        "generic_error": "❌ An error occurred. Please try again.",
        "balance_deduction_failed": "❌ Failed to deduct balance. Please try again.",
        "batch_ready": "✅ *Batch {batch_num}/{total_batches} Ready!*\n\n📚 *{batch_name} - Part {batch_num}/{total_batches}*\n📝 *Questions:* {num_questions}\n⏱️ *Time per question:* {time} seconds\n👤 *Quiz info:* {author}\n\nStart solving or share with your group!",
        "start_quiz": "🎯 Start Quiz",
        "share_quiz": "📤 Share Quiz",
        "share_to_group": "👥 Share to Group",
        "download_word": "📄 Download as Word",
        "confirm_button": "✅ Confirm",
        "cancel_button": "❌ Cancel",
        "topup_balance": "💰 Top Up Balance",
        "insufficient_balance": "❌ *Insufficient Balance*\n\n💰 Required: {price:,} UZS\n💳 Your balance: {balance:,} UZS\n📊 Total quizzes: {total}\n\nPlease top up your balance to continue.",

        "split_ask_randomize": "🔀 *Randomization Options*\n\nChoose how to randomize your quiz:\n\n• *Overall Random*: Both questions and answers randomized\n• *Sequential*: Original order (no randomization)\n• *Questions Only*: Randomize question order, keep answer order\n• *Answers Only*: Keep question order, randomize answers",
        "randomize_overall": "🔀 Overall Random",
        "randomize_sequential": "📝 Keep Sequential",
        "randomize_questions_only": "🔢 Randomize Questions Only",
        "randomize_answers_only": "🔄 Randomize Answers Only",
        "full_batch_option": "📦 All Questions (1 batch)",
                "edit_quiz_info": "✏️ Edit Quiz Info",



    },
    "uz": {
        "split_quiz_start": "📂 *Hemis Faylini Bo'lish*\n\nIltimos, quiz faylingizni yuklang (.txt, .docx, .doc yoki .pdf).\n\n*Qo'llab-quvvatlanadigan formatlar:*\n• ==== va ++++ bilan ajratilgan savollar\n• +/- belgilari bilan savollar\n• Raqamlangan savollar (1., 2., 3.)\n• #, ✅ yoki + bilan to'g'ri javoblar\n\n Iltimos, to'g'ri va xato javoblar aniq belgilanganiga ishonch hosil qiling!",
        "split_file_received": "✅ *Fayl qabul qilindi!*\n\n📊 Jami testlar topildi: *{total}*\n\nEndi har bir to'plamda nechta savol bo'lishini tanlang:",
        "split_invalid_file": "❌ *Noto'g'ri fayl formati* yoki to'g'ri testlar topilmadi.\n\nIltimos tekshiring:\n• Fayl .txt, .docx, .doc yoki .pdf formatida\n• Savollar kamida 2 ta javobga ega\n• To'g'ri javoblar belgilangan",
        "split_choose_time": "⏱ *Har Bir Test Uchun Vaqt*\n\nHar bir test uchun vaqt chegarasini tanlang (soniyalarda):",
        "split_confirm": "✅ *Tasdiqlash*\n\n📊 *Xulosa:*\n• Jami testlar: *{total}*\n• To'plam hajmi: *{batch_size}*\n• To'plamlar soni: *{num_batches}*\n• Oxirgi to'plam: *{last_batch}* ta test\n• Har bir test uchun: *{time}* soniya",
        "split_success": "✅ *Quiz Bo'lish Tugadi!*\n\n{num_batches} ta to'plam muvaffaqiyatli yaratildi.",
        "split_cancelled": "❌ Quiz bo'lish bekor qilindi. Qayta boshlash uchun /start.",
        "back_to_menu": "↩️ Menyuga Qaytish",
        "extracting_text": "📝 {file_type}dan matn ajratilmoqda...",

        "split_ask_quiz_info": "📝 *Quiz Ma'lumoti*\n\nQuiz to'plami haqida ma'lumot kiriting (nomi, muallif, kanal va h.k.).\n\n*Misollar:*\n• Sotsiologiya testlari - Bahtiyor Komilov\n• Filosofiya bo'yicha quizlar - @SimpleLearnuz\n• Tarix savollari to'plami",
        "split_quiz_info_received": "✅ Quiz ma'lumoti saqlandi: *{info}*",
        "quiz_info_too_short": "❌ Quiz ma'lumoti kamida 3 belgidan iborat bo'lishi kerak:",

        "split_batch_name_received": "✅ To'plam nomi o'rnatildi: *{name}*",
        "split_author_received": "✅ Muallif o'rnatildi: *{author}*",

        "could_not_extract": "❌ Fayldan quizlarni ajratib bo'lmadi.\n\n📋 *Qo'llab-quvvatlanadigan formatlar. Iltimos faylingiz formati quyidagilardan birida ekanligiga ishonch hosil qiling.*\n\n*1. Plus/Minus formati:*\n```\n#Savol matni\n+To'g'ri javob\n-Noto'g'ri javob\n-Noto'g'ri javob\n```\n\n*2. Raqamlangan format:*\n```\n1. Savol matni\nA) #To'g'ri javob\nB) Noto'g'ri javob\nC) Noto'g'ri javob\n```\n\nIltimos, faylingiz ushbu formatlardan biriga mos kelishini tekshiring.",
        "price_label": "💰 *Narx: {price:,} UZS*",
        "doc_error": "❌ .doc faylni o'qishda xatolik. .docx yoki .txt formatida saqlang.\n\nXatolik: {error}",
        "generic_error": "❌ Xatolik yuz berdi. Qayta urinib ko'ring.",
        "balance_deduction_failed": "❌ Balansda yetarli mablag' mavjud emas. /start bosing va balansni to'ldiring.",
        "batch_ready": "✅ *{batch_num}/{total_batches} To'plam Tayyor!*\n\n📚 *{batch_name} - {batch_num}/{total_batches} Qism*\n📝 *Savollar:* {num_questions}\n⏱️ *Har bir savol:* {time} soniya\n👤 *Quiz info:* {author}\n\nTestni boshlang yoki ulashing!",
        "start_quiz": "🎯 Testni Boshlash",
        "share_quiz": "📤 Testni Ulashish",
        "share_to_group": "👥 Guruhga Ulashish",
        "download_word": "📄 Word Yuklab Olish",
        "confirm_button": "✅ Tasdiqlash",
        "cancel_button": "❌ Bekor qilish",
        "topup_balance": "💰 Balansni To'ldirish",
        "insufficient_balance": "❌ *Balansingiz yetarli emas*\n\n💰 Kerak: {price:,} UZS\n💳 Sizning balansingiz: {balance:,} UZS\n📊 Fayldagi jami testlar: {total}\n\nBalansni to'ldiring.",

"split_ask_randomize": "🔀 *Aralash Tartiblash Variantlari*\n\nQuizni qanday aralashtirishni tanlang:\n\n• *Hammasini aralashtirish*: Savollar va javoblar aralash bo'ladi\n• *Ketma-ket*: Savollar va variantlarning asl tartibi saqlanadi\n• *Faqat savollar*: Faqat savollar tartibi aralash (21, 14, 2, 10), javoblar (A, B, C, D) o'z tartibida\n• *Faqat javoblar*: Savollar o'z tartibida (1,2,3,4...), faqat javoblar (B, D, C, A) aralash holda bo'ladi",
"randomize_overall": "🔀 Hammasini Aralashtirish",
"randomize_sequential": "📝 Ketma-ketlikni Saqlash",
"randomize_questions_only": "🔢 Faqat Savollarni aralashtirish",
"randomize_answers_only": "🔄 Faqat Javoblarni aralashtirish",
        "full_batch_option": "📦 Barcha Savollar (1 to'plam)",
'edit_quiz_info': "✏️ Quiz Ma'lumotlarini Tahrirlash",


    },
    "ru": {
        "split_quiz_start": "📂 *Разделить Файл*\n\nЗагрузите файл с тестами (.txt, .docx, .doc или .pdf).\n\n*Поддерживаемые форматы:*\n• Вопросы разделенные ==== и ++++\n• Вопросы с +/- маркерами\n• Нумерованные вопросы (1., 2., 3.)\n• Формат с #, ✅ или + для правильных ответов",
        "split_file_received": "✅ *Файл получен!*\n\n📊 Всего тестов найдено: *{total}*\n\nВыберите количество тестов в наборе:",
        "split_invalid_file": "❌ *Неверный формат* или тесты не найдены.\n\nУбедитесь:\n• Файл .txt, .docx, .doc или .pdf\n• Вопросы имеют минимум 2 варианта\n• Правильные ответы отмечены",
        "split_choose_time": "⏱ *Время на Тест*\n\nВыберите время на каждый тест (в секундах):",
        "split_confirm": "✅ *Подтверждение*\n\n📊 *Итого:*\n• Всего тестов: *{total}*\n• Размер набора: *{batch_size}*\n• Количество наборов: *{num_batches}*\n• Последний набор: *{last_batch}* тестов\n• Время на тест: *{time}* секунд",
        "split_success": "✅ *Разделение Завершено!*\n\nСоздано {num_batches} набор(ов).",
        "split_cancelled": "❌ Разделение отменено. Используйте /start.",
        "back_to_menu": "↩️ Вернуться",
        "extracting_text": "📝 Извлечение текста из {file_type}...",
      
        "split_ask_quiz_info": "📝 *Информация о Квизе*\n\nВведите информацию об этом наборе тестов (название, автор, канал и т.д.).\n\n*Примеры:*\n• Тесты по социологии - Бахтиёр Комилов\n• Квизы по философии - @SimpleLearnuz\n• Сборник вопросов по истории",
        "split_quiz_info_received": "✅ Информация сохранена: *{info}*",
        "quiz_info_too_short": "❌ Информация должна быть не менее 3 символов:",
     
        "split_batch_name_received": "✅ Название установлено: *{name}*",
        "split_author_received": "✅ Автор установлен: *{author}*",
        "batch_name_too_short": "❌ Название должно быть не менее 3 символов:",
        "author_name_too_short": "❌ Имя должно быть не менее 2 символов:",
            "could_not_extract": "❌ Не удалось извлечь текст из файла.\n\n📋 *Поддерживаемые Форматы:*\n\n*1. Формат с Разделителями:*\n```\nТекст вопроса====# Правильный вариант====Неправильный вариант====Неправильный вариант++++\n```\n\n*2. Формат Плюс/Минус:*\n```\n#Текст вопроса\n+Правильный вариант\n-Неправильный вариант\n-Неправильный вариант\n```\n\n*3. Нумерованный Формат:*\n```\n1. Текст вопроса\nA) #Правильный вариант\nB) Неправильный вариант\nC) Неправильный вариант\n```\n\nПожалуйста, убедитесь, что ваш файл соответствует одному из этих форматов.",        
            
            
        "insufficient_balance": "❌ *Balansingiz yetarli emas*\n\n💰 Kerak: {price:,} UZS\n💳 Sizning balansingiz: {balance:,} UZS\n📊 Jami testlar: {total}\n\nBalansni to'ldiring.",
        "insufficient_balance": "❌ *Недостаточно средств*\n\n💰 Необходимо: {price:,} UZS\n💳 Баланс: {balance:,} UZS\n📊 Всего тестов: {total}\n\nПополните баланс.",
        "price_label": "💰 *Цена: {price:,} UZS*",
        "doc_error": "❌ Ошибка чтения .doc. Сохраните как .docx или .txt.\n\nОшибка: {error}",
        "generic_error": "❌ Произошла ошибка. Попробуйте снова.",
        "balance_deduction_failed": "❌ Не удалось списать средства.",
        "batch_ready": "✅ *Набор {batch_num}/{total_batches} Готов!*\n\n📚 *{batch_name} - Часть {batch_num}/{total_batches}*\n📝 *Вопросов:* {num_questions}\n⏱️ *Время:* {time} секунд\n👤 *Квиз инфо:* {author}\n\nНачните или поделитесь!",
        "start_quiz": "🎯 Начать Тест",
        "share_quiz": "📤 Поделиться",
        "share_to_group": "👥 В Группу",
        "download_word": "📄 Скачать Word",
        "confirm_button": "✅ Подтвердить",
        "cancel_button": "❌ Отмена",
        "topup_balance": "💰 Пополнить",
        "insufficient_balance": "❌ *Недостаточно средств*\n\n💰 Необходимо: {price:,} UZS\n💳 Баланс: {balance:,} UZS\n📊 Всего тестов: {total}\n\nПополните баланс.",
"split_ask_randomize": "🔀 *Варианты Рандомизации*\n\nВыберите способ перемешивания:\n\n• *Полностью случайно*: Вопросы и ответы перемешаны\n• *Последовательно*: Исходный порядок (без перемешивания)\n• *Только вопросы*: Перемешать порядок вопросов, ответы по порядку\n• *Только ответы*: Вопросы по порядку, ответы перемешаны",
"randomize_overall": "🔀 Полностью Случайно",
"randomize_sequential": "📝 Сохранить Порядок",
"randomize_questions_only": "🔢 Только Вопросы",
"randomize_answers_only": "🔄 Только Ответы",
        "full_batch_option": "📦 Все Вопросы (1 набор)",
'edit_quiz_info': "✏️ Редактировать Информацию",

    }
}

def is_free_user(user_id) -> bool:
    """IDs may be stored as str or int - compare as strings."""
    uid = str(user_id).strip()
    try:
        from src.imports import FREE_USER_IDS
        if uid in {str(x).strip() for x in FREE_USER_IDS}:
            return True
    except Exception:
        pass
    try:
        from src.imports import is_free_user as _f
        return bool(_f(user_id) or _f(uid))
    except Exception:
        return False


def get_translation(language: str, key: str, **kwargs) -> str:
    """Get translation with proper fallback."""
    lang_dict = SPLIT_TRANSLATIONS.get(language, SPLIT_TRANSLATIONS["en"])
    translation = lang_dict.get(key, SPLIT_TRANSLATIONS["en"].get(key, key))
    
    if kwargs:
        try:
            return translation.format(**kwargs)
        except KeyError as e:
            logger.error(f"Missing format key {e} for '{key}' in '{language}'")
            return translation
    return translation


async def start_split_quiz(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Start the quiz split process."""
    from src.user_management import get_user_language
    
    user = update.effective_user
    language = await get_user_language(user.id) or 'en'
    
    # Clear any existing split quiz data
    keys_to_clear = [
        'split_quizzes', 'split_file_path', 'split_batch_size', 
        'split_time', 'split_batch_name', 'split_author', 'split_price'
    ]
    for key in keys_to_clear:
        context.user_data.pop(key, None)
    
    context.user_data['language'] = language
    
    logger.info(f"[SPLIT] User {user.id} started with language: {language}")
    
    text = get_translation(language, "split_quiz_start")
    keyboard = [[KeyboardButton(get_translation(language, "back_to_menu"))]]
    reply_markup = ReplyKeyboardMarkup(keyboard, resize_keyboard=True)
    
    # Handle both message and callback query
    if update.callback_query:
        await update.callback_query.message.reply_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
    else:
        await update.message.reply_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
    
    return SPLIT_WAITING_FILE


async def handle_quiz_file(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle uploaded quiz file."""
    language = context.user_data.get('language', 'en')
    
    if not update.message.document:
        await update.message.reply_text(
            get_translation(language, "could_not_extract"),
            parse_mode=ParseMode.MARKDOWN
        )
        return SPLIT_WAITING_FILE
    
    document = update.message.document
    file_name = document.file_name.lower()
    
    allowed_extensions = ('.txt', '.docx', '.doc', '.pdf', '.xlsx', '.xls')
    if not file_name.endswith(allowed_extensions):
        await update.message.reply_text(
            get_translation(language, "could_not_extract"),
            parse_mode=ParseMode.MARKDOWN
        )
        return SPLIT_WAITING_FILE
    
    file_path = None
    status_message = None
    
    try:
        file_type = os.path.splitext(file_name)[1].upper().replace('.', '')
        status_message = await update.message.reply_text(
            get_translation(language, "extracting_text", file_type=file_type),
            parse_mode=ParseMode.MARKDOWN
        )
        
        file = await context.bot.get_file(document.file_id)
        file_path = f"temp_quiz_{update.effective_user.id}{os.path.splitext(file_name)[1]}"
        await file.download_to_drive(file_path)
        
        content = ""
        
        if file_name.endswith('.txt'):
            # Try multiple encodings for Uzbek/Cyrillic text
            encodings_to_try = ['utf-8', 'windows-1251', 'cp1251', 'utf-16', 'cp1252', 'latin-1']
            content = None
            successful_encoding = None
            
            for encoding in encodings_to_try:
                try:
                    with open(file_path, 'r', encoding=encoding, errors='strict') as f:
                        content = f.read()
                    successful_encoding = encoding
                    logger.info(f"[SPLIT] Successfully read .txt file with {encoding} encoding")
                    break
                except (UnicodeDecodeError, UnicodeError):
                    continue
                except Exception as e:
                    logger.error(f"[SPLIT] Unexpected error with {encoding}: {e}")
                    continue
            
            if content is None:
                # Last resort: read as binary and decode with errors='replace'
                try:
                    with open(file_path, 'rb') as f:
                        raw_content = f.read()
                    # Try to detect encoding or use replace
                    content = raw_content.decode('windows-1251', errors='replace')
                    logger.warning("[SPLIT] Read .txt file with errors='replace' fallback using windows-1251")
                except Exception as e:
                    logger.error(f"[SPLIT] Failed all encoding attempts: {e}")
                    await update.message.reply_text(
                        get_translation(language, "could_not_extract"),
                        parse_mode=ParseMode.MARKDOWN
                    )
                    if file_path and os.path.exists(file_path):
                        os.remove(file_path)
                    return SPLIT_WAITING_FILE
        
        elif file_name.endswith('.docx'):
            doc = Document(file_path)
            content_parts = []
            
            for p in doc.paragraphs:
                if p.text.strip():
                    content_parts.append(p.text.strip())
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text and len(cell_text) > 1:
                            if cell_text not in ['====', '++++', '----', '===', '+++', '---']:
                                content_parts.append(cell_text)
            
            content = '\n'.join(content_parts) 

        elif file_name.endswith('.doc'):
            try:
                ole = olefile.OleFileIO(file_path)
                text_parts = []
                
                for stream_name in ['WordDocument', '1Table', '0Table']:
                    if ole.exists(stream_name):
                        try:
                            stream = ole.openstream(stream_name)
                            raw_data = stream.read()
                            
                            for encoding in ['utf-16-le', 'utf-8', 'cp1251', 'cp1252', 'latin-1']:
                                try:
                                    decoded = raw_data.decode(encoding, errors='ignore')
                                    cleaned = ''.join(c for c in decoded if c.isprintable() or c in '\n\r\t ')
                                    lines = [line.strip() for line in cleaned.split('\n') if len(line.strip()) > 2]
                                    if len(lines) > 5:
                                        text_parts.extend(lines)
                                        break
                                except:
                                    continue
                        except:
                            continue
                
                ole.close()
                content = '\n'.join(text_parts)
                
                if not content.strip() or len(content.strip()) < 50:
                    raise ValueError("Could not extract meaningful text from .doc file")
                    
            except Exception as e:
                logger.error(f"Error reading .doc: {e}")
                if status_message:
                    try:
                        await status_message.delete()
                    except:
                        pass
                
                error_msg = get_translation(language, "doc_error", error=str(e)[:100])
                error_msg += "\n\n💡 *Tip:* Please save your file as .docx or .txt format for better compatibility."
                
                await update.message.reply_text(
                    error_msg,
                    parse_mode=ParseMode.MARKDOWN
                )
                if file_path and os.path.exists(file_path):
                    os.remove(file_path)
                return SPLIT_WAITING_FILE
            
        elif file_name.endswith('.pdf'):
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        content += page_text + '\n'
        
        if not content.strip():
            if status_message:
                try:
                    await status_message.delete()
                except:
                    pass
            await update.message.reply_text(
                get_translation(language, "could_not_extract"),
                parse_mode=ParseMode.MARKDOWN
            )
            if file_path and os.path.exists(file_path):
                os.remove(file_path)
            return SPLIT_WAITING_FILE
        
        logger.info(f"[SPLIT] Content length: {len(content)}")
        
        correct_marker = context.user_data.get('correct_marker', '#')
        quizzes = QuizSplitter.parse_quiz_file(content, correct_marker)
        
        logger.info(f"[SPLIT] Parsed {len(quizzes)} quizzes")
        
        valid_quizzes = []
        issues = []

        for idx, quiz in enumerate(quizzes, 1):
            has_question = quiz.get('question') and len(quiz.get('question', '').strip()) >= 3
            has_options = quiz.get('options') and len(quiz['options']) >= 2
            
            # Handle both single and multiple correct answers
            has_single_correct = quiz.get('correct_answer') is not None
            has_multiple_correct = quiz.get('correct_answers') is not None and len(quiz.get('correct_answers', [])) > 0
            has_correct = has_single_correct or has_multiple_correct
            
            # Validate correct answer indices are in range
            if has_single_correct:
                correct_in_range = quiz['correct_answer'] < len(quiz.get('options', []))
            elif has_multiple_correct:
                correct_in_range = all(idx < len(quiz.get('options', [])) for idx in quiz['correct_answers'])
            else:
                correct_in_range = False
            
            valid_options = [opt for opt in quiz.get('options', []) if opt and len(opt.strip()) > 0]
            has_valid_options = len(valid_options) >= 2
            
            if has_question and has_options and has_correct and correct_in_range and has_valid_options:
                quiz['options'] = valid_options
                # Adjust indices if options were filtered
                if has_single_correct and quiz['correct_answer'] >= len(valid_options):
                    quiz['correct_answer'] = 0
                elif has_multiple_correct:
                    # Filter out any indices that are now out of range
                    quiz['correct_answers'] = [idx for idx in quiz['correct_answers'] if idx < len(valid_options)]
                    if not quiz['correct_answers']:
                        quiz['correct_answers'] = [0]
                
                valid_quizzes.append(quiz)
            else:
                issue_parts = []
                if not has_question:
                    issue_parts.append("no question")
                if not has_valid_options:
                    issue_parts.append("invalid options")
                if not has_correct or not correct_in_range:
                    issue_parts.append("no correct answer")
                issues.append(f"Quiz {idx}: {', '.join(issue_parts)}")

        # ADD LOGGING HERE BEFORE OVERWRITING
        logger.info(f"[SPLIT] Validation complete. Valid quizzes: {len(valid_quizzes)}, Invalid: {len(issues)}")

        quizzes = valid_quizzes  # This was overwriting the count

        # ADD LOGGING HERE AFTER OVERWRITING
        logger.info(f"[SPLIT] After validation - Total quizzes to process: {len(quizzes)}")
                
        if len(quizzes) == 0:
            if status_message:
                try:
                    await status_message.delete()
                except:
                    pass
            
            # First message: Could not extract
            error_msg = get_translation(language, "could_not_extract")
            if issues:
                error_msg += f"\n\n⚠️ Issues found:\n" + "\n".join(issues[:5])
                if len(issues) > 5:
                    error_msg += f"\n... and {len(issues) - 5} more"
            
            await update.message.reply_text(error_msg, parse_mode=ParseMode.MARKDOWN)
            
            # Second message: Helpful tip (separate message)
            helpful_tip = {
                'en': "💡 <b>Tip:</b> You can also create a quiz from this file using /start → 🎯 Create Quiz. Or send you file to @SimpleLearn_main_admin. We will check and help you!",
                'uz': "💡 <b>Maslahat:</b> Siz ushbu fayldan /start → 🎯 Quiz yaratish orqali ham viktorina yaratishingiz mumkin. Yoki faylni @SimpleLearn_main_admin ga yuboring. Biz tekshirib, yordam beramiz!",
                'ru': "💡 <b>Совет:</b> Вы также можете создать викторину из этого файла через /start → 🎯 Создать викторину. Или отправьте файл @SimpleLearn_main_admin. Мы проверим и поможем!"
            }
            
            await update.message.reply_text(
                helpful_tip.get(language, helpful_tip['en']),
                parse_mode=ParseMode.HTML
            )
            
            if file_path and os.path.exists(file_path):
                os.remove(file_path)
            context.user_data.clear()
            return ConversationHandler.END 
        
        # ✅ CHECK BALANCE IMMEDIATELY AFTER PARSING
        total_quizzes = len(quizzes)
        split_price = get_split_pricing(total_quizzes)
        
        from src.user_management import get_user_balance
        user_balance = await get_user_balance(update.effective_user.id)
        
        if not is_free_user(update.effective_user.id) and user_balance < split_price * 100:  # Convert to minor units for comparison
            if status_message:
                try:
                    await status_message.delete()
                except:
                    pass
            
            balance_msg = get_translation(
                language,
                "insufficient_balance",
                price=split_price,
                balance=user_balance // 100,  # Convert back to major units for display
                total=total_quizzes
            )
            
            keyboard = [
                [InlineKeyboardButton(
                    get_translation(language, "topup_balance"),
                    callback_data="topup_balance"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "cancel_button"),
                    callback_data="split_cancel"
                )]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await update.message.reply_text(
                balance_msg,
                parse_mode=ParseMode.MARKDOWN,
                reply_markup=reply_markup
            )
            
            if file_path and os.path.exists(file_path):
                os.remove(file_path)
            return SPLIT_WAITING_FILE
            
        if status_message:
            try:
                await status_message.delete()
            except:
                pass
        
        # Store data for later use
        context.user_data['split_quizzes'] = quizzes
        context.user_data['split_file_path'] = file_path
        context.user_data['split_price'] = split_price
        
        total = len(quizzes)
        text = get_translation(language, "split_file_received", total=total)
        if not is_free_user(update.effective_user.id):
            text += "\n\n" + get_translation(language, "price_label", price=split_price)

        options = QuizSplitter.get_split_options()
        keyboard = []
        row = []

        logger.info(f"[SPLIT] Creating buttons for {total} quizzes with options: {options}")

        # Full batch option text based on language
        full_batch_texts = {
            'en': f"📦 All Questions (1 batch)",
            'uz': f"📦 Barcha savollar (1 to'plam)",
            'ru': f"📦 Все вопросы (1 набор)"
        }
        
        keyboard.append([InlineKeyboardButton(
            full_batch_texts.get(language, full_batch_texts['en']),
            callback_data=f"split_size_{total}"
        )])

        # If total quizzes are less than 20, show only one "All questions" button
        if total >= 20:
            for size in options:
                if size < total:
                    num_batches, _ = QuizSplitter.calculate_batches(total, size)
                    
                    # Format button text based on language
                    if language == 'uz':
                        if num_batches == 1:
                            button_text = f"{size} ta (1 to'plam)"
                        else:
                            button_text = f"{size} ta ({num_batches} to'plam)"
                    elif language == 'ru':
                        if num_batches == 1:
                            button_text = f"{size} (1 набор)"
                        else:
                            button_text = f"{size} ({num_batches} наборов)"
                    else:  # English
                        if num_batches == 1:
                            button_text = f"{size} (1 batch)"
                        else:
                            button_text = f"{size} ({num_batches} batches)"
                    
                    row.append(InlineKeyboardButton(
                        button_text,
                        callback_data=f"split_size_{size}"
                    ))
                    
                    if len(row) == 2:
                        keyboard.append(row)
                        row = []

            # Add remaining button if exists
            if row:
                keyboard.append(row)

            logger.info(f"[SPLIT] Created {len(keyboard)} rows of buttons with {sum(len(r) for r in keyboard)} total buttons")

        # Add cancel button
        keyboard.append([InlineKeyboardButton(
            get_translation(language, "cancel_button"),
            callback_data="split_cancel"
        )])
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await update.message.reply_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
        
        return SPLIT_CHOOSE_SIZE
            
    # At every error return point, add:
    except Exception as e:
        logger.error(f"Error processing file: {e}")
        import traceback
        logger.error(traceback.format_exc())
        
        if status_message:
            try:
                await status_message.delete()
            except:
                pass
        
        await update.message.reply_text(
            get_translation(language, "generic_error"),
            parse_mode=ParseMode.MARKDOWN
        )
        # FIX: Clear context and end conversation on error
        context.user_data.clear()
        return ConversationHandler.END
    
    finally:
        if file_path and os.path.exists(file_path):
            try:
                os.remove(file_path)
                logger.info(f"[SPLIT] Cleaned up temp file: {file_path}")
            except Exception as e:
                logger.error(f"[SPLIT] Failed to delete temp file: {e}")

async def handle_quiz_text(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle quiz text sent as a message (not as file)."""
    language = context.user_data.get('language', 'en')
    
    try:
        content = update.message.text.strip()
        
        if not content or len(content) < 50:
            await update.message.reply_text(
                get_translation(language, "could_not_extract"),
                parse_mode=ParseMode.MARKDOWN
            )
            context.user_data.clear()
            return ConversationHandler.END
        
        status_message = await update.message.reply_text(
            "📝 Processing your quiz text...",
            parse_mode=ParseMode.MARKDOWN
        )
        
        logger.info(f"[SPLIT] Processing text message, length: {len(content)}")
        
        correct_marker = context.user_data.get('correct_marker', '#')
        quizzes = QuizSplitter.parse_quiz_file(content, correct_marker)
        
        logger.info(f"[SPLIT] Parsed {len(quizzes)} quizzes from text")
        
        # Validate quizzes
        valid_quizzes = []
        issues = []

        for idx, quiz in enumerate(quizzes, 1):
            has_question = quiz.get('question') and len(quiz.get('question', '').strip()) >= 3
            has_options = quiz.get('options') and len(quiz['options']) >= 2
            
            has_single_correct = quiz.get('correct_answer') is not None
            has_multiple_correct = quiz.get('correct_answers') is not None and len(quiz.get('correct_answers', [])) > 0
            has_correct = has_single_correct or has_multiple_correct
            
            if has_single_correct:
                correct_in_range = quiz['correct_answer'] < len(quiz.get('options', []))
            elif has_multiple_correct:
                correct_in_range = all(idx < len(quiz.get('options', [])) for idx in quiz['correct_answers'])
            else:
                correct_in_range = False
            
            valid_options = [opt for opt in quiz.get('options', []) if opt and len(opt.strip()) > 0]
            has_valid_options = len(valid_options) >= 2
            
            if has_question and has_options and has_correct and correct_in_range and has_valid_options:
                quiz['options'] = valid_options
                if has_single_correct and quiz['correct_answer'] >= len(valid_options):
                    quiz['correct_answer'] = 0
                elif has_multiple_correct:
                    quiz['correct_answers'] = [idx for idx in quiz['correct_answers'] if idx < len(valid_options)]
                    if not quiz['correct_answers']:
                        quiz['correct_answers'] = [0]
                
                valid_quizzes.append(quiz)
            else:
                issue_parts = []
                if not has_question:
                    issue_parts.append("no question")
                if not has_valid_options:
                    issue_parts.append("invalid options")
                if not has_correct or not correct_in_range:
                    issue_parts.append("no correct answer")
                issues.append(f"Quiz {idx}: {', '.join(issue_parts)}")

        logger.info(f"[SPLIT] Validation complete. Valid quizzes: {len(valid_quizzes)}, Invalid: {len(issues)}")

        quizzes = valid_quizzes
        
        if len(quizzes) == 0:
            if status_message:
                try:
                    await status_message.delete()
                except:
                    pass
            
            error_msg = get_translation(language, "could_not_extract")
            if issues:
                error_msg += f"\n\n⚠️ Issues found:\n" + "\n".join(issues[:5])
                if len(issues) > 5:
                    error_msg += f"\n... and {len(issues) - 5} more"
            
            await update.message.reply_text(error_msg, parse_mode=ParseMode.MARKDOWN)
            
            helpful_tip = {
                'en': "💡 <b>Tip:</b> You can also create a quiz from this text using /start → 🎯 Create Quiz.",
                'uz': "💡 <b>Maslahat:</b> Siz ushbu matndan /start → 🎯 Quiz yaratish orqali ham viktorina yaratishingiz mumkin.",
                'ru': "💡 <b>Совет:</b> Вы также можете создать викторину из этого текста через /start → 🎯 Создать викторину."
            }
            
            await update.message.reply_text(
                helpful_tip.get(language, helpful_tip['en']),
                parse_mode=ParseMode.HTML
            )
            
            context.user_data.clear()
            return ConversationHandler.END 
        
        # Check balance
        total_quizzes = len(quizzes)
        split_price = get_split_pricing(total_quizzes)
        
        from src.user_management import get_user_balance
        user_balance = await get_user_balance(update.effective_user.id)
        
        if not is_free_user(update.effective_user.id) and user_balance < split_price * 100:
            if status_message:
                try:
                    await status_message.delete()
                except:
                    pass
            
            balance_msg = get_translation(
                language,
                "insufficient_balance",
                price=split_price,
                balance=user_balance // 100,
                total=total_quizzes
            )
            
            keyboard = [
                [InlineKeyboardButton(
                    get_translation(language, "topup_balance"),
                    callback_data="topup_balance"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "cancel_button"),
                    callback_data="split_cancel"
                )]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await update.message.reply_text(
                balance_msg,
                parse_mode=ParseMode.MARKDOWN,
                reply_markup=reply_markup
            )
            
            return SPLIT_WAITING_FILE
        
        if status_message:
            try:
                await status_message.delete()
            except:
                pass
        
        # Store data
        context.user_data['split_quizzes'] = quizzes
        context.user_data['split_file_path'] = None  # No file for text input
        context.user_data['split_price'] = split_price
        
        total = len(quizzes)
        text = get_translation(language, "split_file_received", total=total)
        if not is_free_user(update.effective_user.id):
            text += "\n\n" + get_translation(language, "price_label", price=split_price)

        options = QuizSplitter.get_split_options()
        keyboard = []

        keyboard.append([InlineKeyboardButton(
            get_translation(language, "full_batch_option"),
            callback_data=f"split_size_{total}"
        )])

        if total >= 20:
            row = []
            for size in options:
                if size < total:
                    num_batches, _ = QuizSplitter.calculate_batches(total, size)
                    
                    if num_batches == 1:
                        button_text = f"{size} (1 batch)"
                    else:
                        button_text = f"{size} ({num_batches} batches)"
                    
                    row.append(InlineKeyboardButton(
                        button_text,
                        callback_data=f"split_size_{size}"
                    ))
                    
                    if len(row) == 2:
                        keyboard.append(row)
                        row = []

            if row:
                keyboard.append(row)

        keyboard.append([InlineKeyboardButton(
            get_translation(language, "cancel_button"),
            callback_data="split_cancel"
        )])
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await update.message.reply_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
        
        return SPLIT_CHOOSE_SIZE
        
    except Exception as e:
        logger.error(f"Error processing text: {e}")
        import traceback
        logger.error(traceback.format_exc())
        
        await update.message.reply_text(
            get_translation(language, "generic_error"),
            parse_mode=ParseMode.MARKDOWN
        )
        return SPLIT_WAITING_FILE

async def choose_batch_size(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle batch size selection."""
    query = update.callback_query
    await query.answer()
    
    language = context.user_data.get('language', 'en')
    
    if query.data == "split_cancel":
        await query.message.edit_text(get_translation(language, "split_cancelled"))
        file_path = context.user_data.get('split_file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        context.user_data.clear()
        return ConversationHandler.END
    
    batch_size = int(query.data.split('_')[-1])
    context.user_data['split_batch_size'] = batch_size
    
    text = get_translation(language, "split_choose_time")
    
    time_options = [15, 30, 45, 60, 90, 120]
    keyboard = []
    row = []
    for time_sec in time_options:
        row.append(InlineKeyboardButton(f"{time_sec}s", callback_data=f"split_time_{time_sec}"))
        if len(row) == 3:
            keyboard.append(row)
            row = []
    
    if row:
        keyboard.append(row)
    
    keyboard.append([InlineKeyboardButton(
        get_translation(language, "cancel_button"),
        callback_data="split_cancel"
    )])
    
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    await query.message.edit_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
    
    return SPLIT_CHOOSE_TIME


async def choose_time(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle time selection."""
    query = update.callback_query
    await query.answer()
    
    language = context.user_data.get('language', 'en')
    
    if query.data == "split_cancel":
        await query.message.edit_text(get_translation(language, "split_cancelled"))
        file_path = context.user_data.get('split_file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        context.user_data.clear()
        return ConversationHandler.END
    
    time_per_quiz = int(query.data.split('_')[-1])
    context.user_data['split_time'] = time_per_quiz
    
    text = get_translation(language, "split_ask_quiz_info")
    
    keyboard = [[InlineKeyboardButton(
        get_translation(language, "cancel_button"),
        callback_data="split_cancel"
    )]]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    await query.message.edit_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
    
    return SPLIT_ASK_QUIZ_INFO


async def receive_quiz_info(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle quiz info input (combined name + author)."""
    language = context.user_data.get('language', 'en')
    
    quiz_info = update.message.text.strip()
    
    if len(quiz_info) < 3:
        await update.message.reply_text(
            get_translation(language, "quiz_info_too_short"),
            parse_mode=ParseMode.MARKDOWN
        )
        return SPLIT_ASK_QUIZ_INFO
    
    if len(quiz_info) > 200:
        quiz_info = quiz_info[:200]
    
    context.user_data['split_quiz_info'] = quiz_info
    
    confirmation = get_translation(language, "split_quiz_info_received", info=quiz_info)
    await update.message.reply_text(confirmation, parse_mode=ParseMode.MARKDOWN)
    
    text = get_translation(language, "split_ask_randomize")

    keyboard = [
        [InlineKeyboardButton(
            get_translation(language, "randomize_overall"),
            callback_data="split_randomize_overall"
        )],
        [InlineKeyboardButton(
            get_translation(language, "randomize_sequential"),
            callback_data="split_randomize_sequential"
        )],
        [InlineKeyboardButton(
            get_translation(language, "randomize_questions_only"),
            callback_data="split_randomize_questions"
        )],
        [InlineKeyboardButton(
            get_translation(language, "randomize_answers_only"),
            callback_data="split_randomize_answers"
        )],
        [InlineKeyboardButton(
            get_translation(language, "cancel_button"),
            callback_data="split_cancel"
        )]
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    await update.message.reply_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
    
    return SPLIT_ASK_RANDOMIZE

async def choose_randomization(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle randomization choice."""
    query = update.callback_query
    await query.answer()
    
    language = context.user_data.get('language', 'en')
    
    if query.data == "split_cancel":
        await query.message.edit_text(get_translation(language, "split_cancelled"))
        file_path = context.user_data.get('split_file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        context.user_data.clear()
        return ConversationHandler.END
    
    # Store randomization preference
    randomization_map = {
        "split_randomize_overall": "overall",
        "split_randomize_sequential": "sequential",
        "split_randomize_questions": "questions_only",
        "split_randomize_answers": "answers_only"
    }
    
    randomize_mode = randomization_map.get(query.data, "sequential")
    context.user_data['split_randomize_mode'] = randomize_mode   
    # Validate required data
    if 'split_quizzes' not in context.user_data:
        await query.message.edit_text(
            get_translation(language, "generic_error"),
            parse_mode=ParseMode.MARKDOWN
        )
        context.user_data.clear()
        return ConversationHandler.END
    
    quizzes = context.user_data['split_quizzes']
    batch_size = context.user_data.get('split_batch_size')
    time_per_quiz = context.user_data.get('split_time')
    quiz_info = context.user_data.get('split_quiz_info', 'Quiz Collection')
    
    if not batch_size or not time_per_quiz:
        await query.message.edit_text(
            get_translation(language, "generic_error"),
            parse_mode=ParseMode.MARKDOWN
        )
        context.user_data.clear()
        return ConversationHandler.END
    
    total = len(quizzes)
    num_batches, last_batch = QuizSplitter.calculate_batches(total, batch_size)
    
    text = get_translation(
        language,
        "split_confirm",
        total=total,
        batch_size=batch_size,
        num_batches=num_batches,
        last_batch=last_batch,
        time=time_per_quiz
    )
    


    randomize_labels = {
        "overall": "🔀 Overall Random",
        "sequential": "📝 Sequential",
        "questions_only": "🔢 Questions Only",
        "answers_only": "🔄 Answers Only"
    }
    
    text += f"\n• Quiz Info: *{quiz_info}*"
    text += f"\n• Randomization: *{randomize_labels[randomize_mode]}*"

    
    keyboard = [
        [
            InlineKeyboardButton(
                get_translation(language, "confirm_button"),
                callback_data="split_confirm_yes"
            ),
            InlineKeyboardButton(
                get_translation(language, "cancel_button"),
                callback_data="split_cancel"
            )
        ]
    ]
    
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    await query.message.edit_text(text, reply_markup=reply_markup, parse_mode=ParseMode.MARKDOWN)
    
    return SPLIT_CONFIRM




async def confirm_split(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Handle split confirmation and create batches."""
    query = update.callback_query
    await query.answer()
    
    language = context.user_data.get('language', 'en')
    
    if query.data == "split_cancel":
        await query.message.edit_text(get_translation(language, "split_cancelled"))
        file_path = context.user_data.get('split_file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        context.user_data.clear()
        return ConversationHandler.END
    
    try:
        quizzes = context.user_data['split_quizzes']
        batch_size = context.user_data['split_batch_size']
        time_per_quiz = context.user_data['split_time']
        quiz_info = context.user_data.get('split_quiz_info', 'Quiz Collection')
        randomize_questions = context.user_data.get('split_randomize_questions', False)
        user_id = update.effective_user.id
        split_price = 0 if is_free_user(user_id) else context.user_data.get('split_price', 0)
        
        from payment_integration_new import deduct_balance
        from src.user_management import get_user_balance
        
        if split_price > 0:
            success = await deduct_balance(user_id, split_price * 100)
        else:
            success = True          # free user: nothing to charge
        if success:
            charged = split_price * 100
        if not success:
            user_balance = await get_user_balance(user_id)
            balance_msg = get_translation(
                language, "insufficient_balance",
                price=split_price,
                balance=user_balance // 100,      # was missing // 100
                total=len(quizzes)
            )
            
            keyboard = [
                [InlineKeyboardButton(
                    get_translation(language, "topup_balance"),
                    callback_data="topup_balance"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "cancel_button"),
                    callback_data="split_cancel"
                )]
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await query.message.edit_text(
                balance_msg,
                parse_mode=ParseMode.MARKDOWN,
                reply_markup=reply_markup
            )
            return SPLIT_CONFIRM
        
        logger.info(f"[SPLIT] Deducted {split_price} UZS from user {user_id}")   
        
        batches = QuizSplitter.split_quizzes(quizzes, batch_size)
        
        bot_username = (await context.bot.get_me()).username
        
        randomize_mode = context.user_data.get('split_randomize_mode', 'sequential')

        quiz_ids = []
        for batch_num, batch in enumerate(batches, 1):
            # Apply randomization based on mode
            if randomize_mode == 'overall':
                # Randomize question order
                batch = batch.copy()
                random.shuffle(batch)
            elif randomize_mode == 'questions_only':
                # Randomize question order only
                batch = batch.copy()
                random.shuffle(batch)
            # Note: answers_only and sequential don't need question shuffling here
            
            randomized_batch = []
            for quiz in batch:
                options = quiz['options'].copy()
                
                # Handle answer randomization based on mode
                should_randomize_answers = randomize_mode in ['overall', 'answers_only']
                
                if quiz.get('allows_multiple_answers', False):
                    if should_randomize_answers:
                        correct_answer_texts = [options[idx] for idx in quiz['correct_answers']]
                        random.shuffle(options)
                        new_correct_indices = [options.index(text) for text in correct_answer_texts]
                    else:
                        new_correct_indices = quiz['correct_answers']
                    
                    randomized_batch.append({
                        'question': quiz['question'],
                        'options': options,
                        'correct_answers': new_correct_indices,
                        'allows_multiple_answers': True
                    })
                else:
                    if should_randomize_answers:
                        correct_answer_text = options[quiz['correct_answer']]
                        random.shuffle(options)
                        new_correct_index = options.index(correct_answer_text)
                    else:
                        new_correct_index = quiz['correct_answer']
                    
                    randomized_batch.append({
                        'question': quiz['question'],
                        'options': options,
                        'correct_answer': new_correct_index,
                        'allows_multiple_answers': False
                    })
            
            quiz_data = {
                "user_id": user_id,
                "title": f"{quiz_info} - Part {batch_num}/{len(batches)}",  # Changed
                "summary": f"{len(randomized_batch)} questions - {time_per_quiz}s per question",
                "questions": randomized_batch,
                "quiz_time": time_per_quiz,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "quiz_info": quiz_info  
            }
            
            from src.user_management import save_shared_quiz
            quiz_id = await save_shared_quiz(quiz_data, user_id)
            quiz_ids.append(quiz_id)
            
            quiz_message = get_translation(
                language,
                "batch_ready",
                batch_num=batch_num,
                total_batches=len(batches),
                batch_name=quiz_info,  
                num_questions=len(randomized_batch),
                time=time_per_quiz,
                author=quiz_info  
            )
            
            keyboard = [
                [InlineKeyboardButton(
                    get_translation(language, "start_quiz"),
                    callback_data=f"start_quiz_{quiz_id}"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "share_quiz"),
                    switch_inline_query=f"quiz_{quiz_id}"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "share_to_group"),
                    url=f"https://t.me/{bot_username}?startgroup=groupquiz_{quiz_id}"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "edit_quiz_info"),
                    callback_data=f"edit_quiz_info_{quiz_id}"
                )],
                [InlineKeyboardButton(
                    get_translation(language, "download_word"),
                    callback_data=f"download_quiz_{quiz_id}"
                )],
            ]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            await context.bot.send_message(
                chat_id=user_id,
                text=quiz_message,
                parse_mode="MARKDOWN",
                reply_markup=reply_markup
            )
            
            await asyncio.sleep(0.5)
        
        text = get_translation(language, "split_success", num_batches=len(batches))
        
        await query.message.edit_text(text, parse_mode=ParseMode.MARKDOWN)
        
        file_path = context.user_data.get('split_file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        
        context.user_data.clear()
        
        return ConversationHandler.END
        
    except Exception as e:
        logger.error(f"Error creating batches: {e}")
        import traceback
        logger.error(traceback.format_exc())
        await query.message.edit_text(
            get_translation(language, "generic_error"),
            parse_mode=ParseMode.MARKDOWN
        )
        return ConversationHandler.END



async def cancel_split(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Cancel the split process."""
    from src.user_management import get_user_language
    
    language = context.user_data.get('language', 'en')
    user_id = update.effective_user.id
    
    # Clean up file if exists
    file_path = context.user_data.get('split_file_path')
    if file_path and os.path.exists(file_path):
        os.remove(file_path)
    
    context.user_data.clear()
    
    # Import and call the start command to show main menu
    from handlers.start_handler import start
    
    await start(update, context)
    
    return ConversationHandler.END


def get_split_pricing(total_quizzes: int) -> int:
    """Calculate pricing for quiz splitting (UZS)."""
    if total_quizzes < 10:
        return 5000
    if total_quizzes <= 50:
        return total_quizzes * 200
    elif total_quizzes <= 100:
        return total_quizzes * 170
    elif total_quizzes <= 150:
        return total_quizzes * 150
    elif total_quizzes <= 200:
        return total_quizzes * 150
    else:
        return total_quizzes * 130


def get_split_quiz_handler():
    """Get the conversation handler for quiz splitting."""
    return ConversationHandler(
        entry_points=[
            MessageHandler(
                filters.Regex("^(📂 Split Quiz Hemis|📂 Hemis quizni bo'lish|📂 Разделить квиз Hemis)$"),
                start_split_quiz
            )
        ],
        states={
            SPLIT_WAITING_FILE: [
                MessageHandler(filters.Document.ALL, handle_quiz_file),
                MessageHandler(
                    filters.TEXT & ~filters.COMMAND & ~filters.Regex(
                        "^(↩️|Back|Назад|Orqaga|📚|⚙️|💰|🔗|🎯|📂|❓)"
                    ),
                    handle_quiz_text
                ),
                MessageHandler(filters.Regex("^(↩️|Back|Назад|Orqaga)"), cancel_split)
            ],
            SPLIT_CHOOSE_SIZE: [
                CallbackQueryHandler(choose_batch_size, pattern="^split_(size|cancel)")
            ],
            SPLIT_CHOOSE_TIME: [
                CallbackQueryHandler(choose_time, pattern="^split_(time|cancel)")
            ],
            SPLIT_ASK_QUIZ_INFO: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, receive_quiz_info),
                CallbackQueryHandler(cancel_split, pattern="^split_cancel")
            ],
            SPLIT_ASK_RANDOMIZE: [
                CallbackQueryHandler(choose_randomization, pattern="^split_(randomize_overall|randomize_sequential|randomize_questions|randomize_answers|cancel)")
            ],
            SPLIT_CONFIRM: [
                CallbackQueryHandler(confirm_split, pattern="^split_(confirm_yes|cancel)")
            ]
        },
        fallbacks=[
            CommandHandler("cancel", cancel_split),
            MessageHandler(filters.Regex("^(❌ Cancel|❌ Бекор қилиш|❌ Отмена)$"), cancel_split)
        ],
        name="split_quiz_conversation",
        persistent=False,
        allow_reentry=True
    )